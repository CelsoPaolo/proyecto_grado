"""Reporte por rango en Excel (RF-96, plan.md 7.11).

Presenta EL MISMO diccionario que produce panel.reportes.calcular(); no
hace consultas (RN-REP-07). Una hoja por seccion: Resumen, Productos,
Usuarios y Servicios.

Los importes van como NUMEROS con formato monetario, no como texto: asi
el Administrador puede sumarlos o graficarlos en Excel sin convertir
nada.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from panel.reportes import estados_con_etiqueta

FORMATO_BS = '"Bs" #,##0.00'
NEGRITA = Font(bold=True)


def _encabezado(hoja, *titulos):
    hoja.append(titulos)
    for celda in hoja[hoja.max_row]:
        celda.font = NEGRITA


def _importe(hoja, columna):
    """Aplica el formato monetario a una columna, salvo el encabezado."""
    for (celda,) in hoja.iter_rows(min_row=2, min_col=columna, max_col=columna):
        celda.number_format = FORMATO_BS


def _anchos(hoja, *anchos):
    for letra, ancho in zip("ABCDEF", anchos):
        hoja.column_dimensions[letra].width = ancho


def libro_del_reporte(reporte):
    """Devuelve los bytes del .xlsx."""
    libro = Workbook()

    resumen = libro.active
    resumen.title = "Resumen"
    resumen.append(["Reporte de operación de PC-Solution"])
    resumen["A1"].font = Font(bold=True, size=13)
    resumen.append(["Desde", reporte["desde"]])
    resumen.append(["Hasta", reporte["hasta"]])
    for fila in (2, 3):
        resumen.cell(row=fila, column=2).number_format = "DD/MM/YYYY"
    resumen.append([])
    _encabezado(resumen, "Concepto", "Cantidad", "Importe")
    resumen.append(["Ventas confirmadas", reporte["ventas"]["cantidad"], reporte["ventas"]["importe"]])
    resumen.append(["Servicios finalizados", reporte["servicios"]["cantidad"], reporte["servicios"]["importe"]])
    # RN-REP-03: el total, siempre junto a su desglose.
    resumen.append(["Ingreso total", None, reporte["ingreso_total"]["importe"]])
    resumen.cell(row=resumen.max_row, column=1).font = NEGRITA
    for fila in range(6, resumen.max_row + 1):
        resumen.cell(row=fila, column=3).number_format = FORMATO_BS
    _anchos(resumen, 28, 12, 16)

    productos = libro.create_sheet("Productos")
    _encabezado(productos, "Código", "Producto", "Unidades", "Importe")
    for p in reporte["productos_mas_vendidos"]:
        productos.append([p["codigo"], p["nombre"], p["unidades"], p["importe"]])
    _importe(productos, 4)
    _anchos(productos, 16, 40, 10, 16)

    usuarios = libro.create_sheet("Usuarios")
    _encabezado(usuarios, "Usuario", "Ventas", "Importe")
    for u in reporte["ventas_por_usuario"]:
        usuarios.append([u["nombre_completo"], u["cantidad"], u["importe"]])
    _importe(usuarios, 3)
    _anchos(usuarios, 32, 10, 16)

    servicios = libro.create_sheet("Servicios")
    _encabezado(servicios, "Estado", "Solicitudes registradas")
    for etiqueta, cantidad in estados_con_etiqueta(reporte):
        servicios.append([etiqueta, cantidad])
    _anchos(servicios, 20, 22)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
