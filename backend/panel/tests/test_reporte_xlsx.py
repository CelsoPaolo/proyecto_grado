"""T-12.6 — Reporte por rango en Excel (RF-96, plan.md 7.11).

El archivo se abre con openpyxl en la prueba: una hoja por seccion, y
los importes son celdas NUMERICAS con formato monetario, no texto.
"""

from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from catalogo.tests.factories import ProductoFactory
from usuarios.tests.factories import AdministradorFactory, VendedorFactory
from ventas.services import registrar_venta

pytestmark = pytest.mark.django_db

URL = "/api/v1/reportes/xlsx"


def _autenticar(usuario):
    cliente = APIClient()
    respuesta = cliente.post(
        "/api/v1/auth/login",
        {"username": usuario.username, "password": "ClaveSegura123"},
        format="json",
    )
    cliente.credentials(HTTP_AUTHORIZATION=f"Bearer {respuesta.data['access']}")
    return cliente


@pytest.fixture
def libro():
    producto = ProductoFactory(nombre="Teclado mecánico", stock=10, precio=Decimal("350.50"))
    registrar_venta(
        usuario=VendedorFactory(nombre_completo="Juan Pérez"),
        cliente_nombre="C",
        cliente_telefono="1",
        items=[{"producto_id": producto.pk, "cantidad": 2}],
    )
    hoy = timezone.localdate().isoformat()
    respuesta = _autenticar(AdministradorFactory()).get(URL, {"desde": hoy, "hasta": hoy})
    assert respuesta.status_code == 200
    assert f"reporte-{hoy}_{hoy}.xlsx" in respuesta["Content-Disposition"]
    return load_workbook(BytesIO(respuesta.content)), producto


def test_una_hoja_por_seccion(libro):
    wb, _ = libro

    assert wb.sheetnames == ["Resumen", "Productos", "Usuarios", "Servicios"]


def test_los_importes_son_numeros_con_formato_monetario(libro):
    wb, producto = libro

    resumen = {fila[0].value: fila for fila in wb["Resumen"].iter_rows() if fila[0].value}
    ventas = resumen["Ventas confirmadas"]
    assert ventas[1].value == 1
    assert isinstance(ventas[2].value, (int, float)) and ventas[2].value == pytest.approx(701.00)
    assert "Bs" in ventas[2].number_format
    assert resumen["Ingreso total"][2].value == pytest.approx(701.00)

    fila = next(wb["Productos"].iter_rows(min_row=2))
    assert [c.value for c in fila[:3]] == [producto.codigo, "Teclado mecánico", 2]
    assert isinstance(fila[3].value, (int, float)) and "Bs" in fila[3].number_format

    usuario = next(wb["Usuarios"].iter_rows(min_row=2))
    assert usuario[0].value == "Juan Pérez" and isinstance(usuario[2].value, (int, float))


def test_los_cuatro_estados_con_su_etiqueta(libro):
    wb, _ = libro

    filas = [(f[0].value, f[1].value) for f in wb["Servicios"].iter_rows(min_row=2)]

    assert [e for e, _ in filas] == ["Registrada", "En atención", "Finalizada", "Cancelada"]
    assert all(isinstance(n, int) for _, n in filas)


def test_solo_el_administrador():
    respuesta = _autenticar(VendedorFactory()).get(URL, {"desde": "2026-10-01", "hasta": "2026-10-01"})

    assert respuesta.status_code == 403
