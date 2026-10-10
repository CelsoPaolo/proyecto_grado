"""T-12.7 — Coherencia de las tres salidas del reporte (RN-REP-07,
RN-REP-09, RNF-26, RNF-28).

1. El JSON, el PDF y el Excel del MISMO rango tienen las mismas cifras.
   Se compara cifra a cifra, no «que salgan»: las tres presentan el
   resultado de reportes.calcular(), y esta prueba es la que lo vigila.
2. El carnet de un cliente no aparece en ninguna de las tres.
3. Un año de datos sinteticos: el reporte y cada exportacion, en menos
   de 5 segundos.
"""

import time
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.db import connection
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from catalogo.tests.factories import CategoriaFactory, ProductoFactory
from clientes.tests.factories import ClienteFactory
from servicios.models import EstadoSolicitud, SolicitudServicio
from servicios.tests.factories import SolicitudFactory
from usuarios.tests.factories import AdministradorFactory, TecnicoFactory, VendedorFactory
from ventas.models import DetalleVenta, EstadoVenta, Venta
from ventas.services import anular_venta, registrar_venta

pytestmark = pytest.mark.django_db

URL = "/api/v1/reportes"
CARNET = "7777777-1Z"


def _autenticar(usuario):
    cliente = APIClient()
    respuesta = cliente.post(
        "/api/v1/auth/login",
        {"username": usuario.username, "password": "ClaveSegura123"},
        format="json",
    )
    cliente.credentials(HTTP_AUTHORIZATION=f"Bearer {respuesta.data['access']}")
    return cliente


def _bs(importe):
    """Como lo imprime el PDF: floatformat:2 con la configuracion regional
    del proyecto (coma decimal)."""
    return f"{Decimal(importe):.2f}".replace(".", ",")


@pytest.fixture
def tres_salidas():
    """Un dia con ventas de dos vendedores, una anulada, un cliente con
    carnet, y servicios en los cuatro estados."""
    ana, juan = VendedorFactory(nombre_completo="Ana Rojas"), VendedorFactory(nombre_completo="Juan Pérez")
    cliente = ClienteFactory(nombre="Cliente con carnet", carnet=CARNET)
    a = ProductoFactory(nombre="Teclado mecánico", stock=50, precio=Decimal("350.50"))
    b = ProductoFactory(nombre="Mouse óptico", stock=50, precio=Decimal("85.25"))

    def vender(usuario, lineas, **cliente_kw):
        return registrar_venta(
            usuario=usuario,
            items=[{"producto_id": p.pk, "cantidad": n} for p, n in lineas],
            **(cliente_kw or {"cliente_nombre": "Ocasional", "cliente_telefono": "1"}),
        )

    vender(ana, [(a, 2), (b, 1)])
    vender(juan, [(b, 4)], cliente_id=cliente.pk)
    anulada = vender(juan, [(a, 9)])
    anular_venta(venta=anulada, usuario=juan, motivo="Devolución")

    tecnico = TecnicoFactory()
    SolicitudFactory(registrada_por=tecnico, estado=EstadoSolicitud.REGISTRADA)
    SolicitudFactory(registrada_por=tecnico, estado=EstadoSolicitud.EN_ATENCION)
    SolicitudFactory(
        registrada_por=tecnico, estado=EstadoSolicitud.FINALIZADA, observaciones="Hecho",
        monto_cobrado=Decimal("120.00"), fecha_cierre=timezone.now(), cerrada_por=tecnico,
    )
    SolicitudFactory(
        registrada_por=tecnico, estado=EstadoSolicitud.CANCELADA, motivo_cancelacion="Desistió",
        monto_cobrado=Decimal("999.00"), fecha_cierre=timezone.now(), cerrada_por=tecnico,
    )

    hoy = timezone.localdate().isoformat()
    admin = _autenticar(AdministradorFactory())
    rango = {"desde": hoy, "hasta": hoy}
    json = admin.get(URL, rango)
    pdf = admin.get(f"{URL}/pdf", rango)
    xlsx = admin.get(f"{URL}/xlsx", rango)
    assert json.status_code == pdf.status_code == xlsx.status_code == 200
    texto_pdf = " ".join(
        "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages).split()
    )
    return json.data, texto_pdf, load_workbook(BytesIO(xlsx.content)), (json, pdf, xlsx)


def _hoja(libro, nombre):
    return [[c.value for c in fila] for fila in libro[nombre].iter_rows()]


def test_las_cifras_del_resumen_son_identicas_en_las_tres(tres_salidas):
    datos, texto, libro, _ = tres_salidas
    resumen = {fila[0]: fila for fila in _hoja(libro, "Resumen") if fila and fila[0]}

    # Comprobacion de cordura: la anulada y la cancelada no suman.
    assert datos["ventas"] == {"cantidad": 2, "importe": "1127.25"}
    assert datos["servicios"] == {"cantidad": 1, "importe": "120.00"}

    for concepto, seccion in (("Ventas confirmadas", "ventas"), ("Servicios finalizados", "servicios")):
        fila = resumen[concepto]
        assert fila[1] == datos[seccion]["cantidad"]
        assert Decimal(str(fila[2])) == Decimal(datos[seccion]["importe"])
        assert f"{concepto} {datos[seccion]['cantidad']} {_bs(datos[seccion]['importe'])}" in texto

    total = datos["ingreso_total"]["importe"]
    assert Decimal(str(resumen["Ingreso total"][2])) == Decimal(total)
    assert f"Ingreso total {_bs(total)}" in texto


def test_los_productos_mas_vendidos_son_identicos_en_las_tres(tres_salidas):
    datos, texto, libro, _ = tres_salidas
    filas = _hoja(libro, "Productos")[1:]

    assert [(f[0], f[1], f[2], Decimal(str(f[3]))) for f in filas] == [
        (p["codigo"], p["nombre"], p["unidades"], Decimal(p["importe"])) for p in datos["productos_mas_vendidos"]
    ]
    for p in datos["productos_mas_vendidos"]:
        assert f"{p['codigo']} {p['nombre']} {p['unidades']} {_bs(p['importe'])}" in texto


def test_las_ventas_por_usuario_son_identicas_en_las_tres(tres_salidas):
    datos, texto, libro, _ = tres_salidas
    filas = _hoja(libro, "Usuarios")[1:]

    assert [(f[0], f[1], Decimal(str(f[2]))) for f in filas] == [
        (u["nombre_completo"], u["cantidad"], Decimal(u["importe"])) for u in datos["ventas_por_usuario"]
    ]
    for u in datos["ventas_por_usuario"]:
        assert f"{u['nombre_completo']} {u['cantidad']} {_bs(u['importe'])}" in texto


def test_los_servicios_por_estado_son_identicos_en_las_tres(tres_salidas):
    datos, texto, libro, _ = tres_salidas
    etiquetas = dict(EstadoSolicitud.choices)
    filas = dict((f[0], f[1]) for f in _hoja(libro, "Servicios")[1:])

    assert datos["servicios_por_estado"] == {"REGISTRADA": 1, "EN_ATENCION": 1, "FINALIZADA": 1, "CANCELADA": 1}
    for valor, cantidad in datos["servicios_por_estado"].items():
        assert filas[etiquetas[valor]] == cantidad
        assert f"{etiquetas[valor]} {cantidad}" in texto


def test_el_carnet_no_aparece_en_ninguna_salida(tres_salidas):
    """RN-REP-09 y RNF-26, buscando el VALOR del carnet de un cliente que
    si compro en el rango, en los bytes y en el texto de cada salida."""
    _, texto, libro, (json, pdf, xlsx) = tres_salidas

    assert CARNET not in json.content.decode()
    assert CARNET not in texto
    celdas = [str(c) for hoja in libro.sheetnames for fila in _hoja(libro, hoja) for c in fila if c is not None]
    assert not any(CARNET in c for c in celdas)
    # El xlsx es un zip: se busca tambien en el XML de dentro.
    import zipfile

    with zipfile.ZipFile(BytesIO(xlsx.content)) as z:
        assert not any(CARNET.encode() in z.read(nombre) for nombre in z.namelist())


# ─────────────────────────────────────────────────────────────────────
# RNF-28 — un año de datos en menos de 5 segundos
# ─────────────────────────────────────────────────────────────────────
VENTAS_POR_DIA = 20
DIAS = 365
LIMITE_SEGUNDOS = 5.0


@pytest.fixture
def un_año_de_datos():
    """~7.300 ventas con 2 lineas cada una y ~1.800 servicios cerrados.

    Se insertan en bloque (bulk_create) y las fechas se reparten despues
    con SQL: pasar por registrar_venta() siete mil veces haria la prueba
    lentisima sin medir nada del reporte.
    """
    categoria = CategoriaFactory()
    productos = [ProductoFactory(categoria=categoria, stock=10**6) for _ in range(40)]
    vendedores = [VendedorFactory() for _ in range(4)]
    tecnico = TecnicoFactory()
    total = VENTAS_POR_DIA * DIAS

    def venta(i):
        # Una de cada 25, anulada, con la trazabilidad que exige la base.
        anulada = i % 25 == 0
        return Venta(
            usuario=vendedores[i % 4],
            cliente_nombre="Cliente",
            cliente_telefono="1",
            estado=EstadoVenta.ANULADA if anulada else EstadoVenta.CONFIRMADA,
            total=Decimal("170.00"),
            **(
                {"motivo_anulacion": "Prueba", "anulada_por": vendedores[0], "fecha_anulacion": timezone.now()}
                if anulada
                else {}
            ),
        )

    ventas = Venta.objects.bulk_create(venta(i) for i in range(total))
    DetalleVenta.objects.bulk_create(
        DetalleVenta(
            venta=v,
            producto=productos[(i + k) % 40],
            producto_nombre="P",
            precio_unitario=Decimal("85.00"),
            cantidad=1,
            subtotal=Decimal("85.00"),
        )
        for i, v in enumerate(ventas)
        for k in (0, 7)
    )
    SolicitudServicio.objects.bulk_create(
        SolicitudServicio(
            cliente_nombre="Cliente", cliente_telefono="1", tipo="MANTENIMIENTO", modalidad="TALLER",
            descripcion="x", estado=EstadoSolicitud.FINALIZADA, observaciones="Hecho",
            monto_cobrado=Decimal("100.00"), fecha_cierre=timezone.now(),
            registrada_por=tecnico, cerrada_por=tecnico,
        )
        for _ in range(DIAS * 5)
    )
    inicio = date(2025, 1, 1)
    with connection.cursor() as cursor:
        cursor.execute(
            "UPDATE ventas_venta SET fecha_hora = %s::date + ((id %% %s) * interval '1 day') + interval '15 hour'",
            [inicio, DIAS],
        )
        cursor.execute(
            "UPDATE servicios_solicitudservicio SET fecha_registro = %s::date + ((id %% %s) * interval '1 day'),"
            " fecha_cierre = %s::date + ((id %% %s) * interval '1 day') + interval '18 hour'",
            [inicio, DIAS, inicio, DIAS],
        )
        # RNF-28 dice «bajo carga normal de trabajo». En una base normal,
        # autovacuum mantiene las estadisticas al dia; dentro de la
        # transaccion de esta prueba no puede, porque no ve filas sin
        # confirmar, y el planificador elige planes para tablas vacias.
        # En la suite completa eso llego a 213 s por peticion. ANALYZE si
        # ve las filas de su propia transaccion y deja la base como
        # estaria en uso real.
        cursor.execute(
            "ANALYZE ventas_venta, ventas_detalleventa, servicios_solicitudservicio, catalogo_producto"
        )
    return inicio, inicio + timedelta(days=DIAS - 1)


@pytest.mark.parametrize("salida", ["", "/pdf", "/xlsx"], ids=["json", "pdf", "xlsx"])
def test_un_año_de_reporte_y_su_exportacion_en_menos_de_5_segundos(un_año_de_datos, salida):
    desde, hasta = un_año_de_datos
    admin = _autenticar(AdministradorFactory())

    inicio = time.perf_counter()
    respuesta = admin.get(f"{URL}{salida}", {"desde": desde.isoformat(), "hasta": hasta.isoformat()})
    segundos = time.perf_counter() - inicio

    assert respuesta.status_code == 200
    assert segundos < LIMITE_SEGUNDOS, f"{salida or 'json'}: {segundos:.2f} s"
    if not salida:
        # Que de verdad recorrio el año: 7.300 ventas menos las anuladas.
        assert respuesta.data["ventas"]["cantidad"] == sum(1 for i in range(VENTAS_POR_DIA * DIAS) if i % 25)
